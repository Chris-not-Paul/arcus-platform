"""Read-only comparison of canonical Master with the pre-edit backup."""
import json
from copy import copy
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
before = openpyxl.load_workbook(root / 'tmp/evidence-review-2026-10-19/MASTER_RESEARCH-before.xlsx')
after = openpyxl.load_workbook(root / 'private-data/raw/MASTER_RESEARCH.xlsx')
plan = json.loads((root / 'config/curation/evidence-review-2026-10-19.json').read_text(encoding='utf-8'))
decisions = {d['id']: d['level'] for d in plan['decisions']}
assert before.sheetnames == after.sheetnames
changes = []
for old, new in zip(before.worksheets, after.worksheets):
    assert old.freeze_panes == new.freeze_panes, old.title
    assert str(old.data_validations) == str(new.data_validations), old.title
    assert list(old.tables) == list(new.tables), old.title
    assert str(old.merged_cells) == str(new.merged_cells), old.title
    for row in old:
        for cell in row:
            updated = new[cell.coordinate]
            assert copy(cell.font) == copy(updated.font) and copy(cell.fill) == copy(updated.fill) and copy(cell.border) == copy(updated.border) and copy(cell.alignment) == copy(updated.alignment) and cell.number_format == updated.number_format, (old.title, cell.coordinate, 'style')
            if cell.value == updated.value:
                continue
            changes.append((old.title, cell.coordinate, cell.value, updated.value))
            if old.title == 'EVENTS':
                assert cell.column == 27
                event_id = old.cell(cell.row, 1).value
                assert cell.value == 'Unspecified' and updated.value == decisions[event_id]
            elif old.title == 'CURATION_LOG':
                assert 2218 <= cell.row <= 2244 and 1 <= cell.column <= 8 and cell.value is None
            elif old.title == 'QA_SUMMARY':
                assert cell.coordinate in {'B2', 'B11', 'B12', 'B13', 'B14', 'A16'}
            else:
                raise AssertionError((old.title, cell.coordinate, 'unexpected change'))
assert len([c for c in changes if c[0] == 'EVENTS']) == 27
assert sum(1 for r in after['CURATION_LOG'].iter_rows(min_row=2218, max_row=2244) if r[1].value in decisions) == 27
print(json.dumps({'passed': True, 'event_changes': 27, 'all_unrelated_cell_values_and_styles_preserved': True, 'validations_and_tables_preserved': True, 'changed_cells': len(changes)}))
